"""Synthetic localisation checks and saved-checkpoint validation comparisons."""

import csv
import json
from pathlib import Path

import cv2
import numpy as np
import pytest
import torch
from openpyxl import load_workbook
from torch.nn import functional as F

from Heatmaps.infer_landmarks_peaks_refinement import (
    ComparisonConfig, discover_checkpoints, error_summary, forward_for_comparison,
    prepare_jobs, run_comparison,
)
from Heatmaps.utils.heatmap_inference_utils import (
    HeatmapImageRecord, build_config_from_checkpoint_metadata, load_model_from_checkpoint,
    run_heatmap_inference_for_records, write_csv_rows,
)
from Heatmaps.utils.io_utils import heatmaps_to_points, scale_points_to_original
from Heatmaps.utils.peak_refinement import PeakRefinementConfig, decode_peak_comparison, refine_peak
import test_inference as inference_fixtures


@pytest.fixture(autouse=True)
def limited_threads():
    previous = torch.get_num_threads()
    torch.set_num_threads(2)
    yield
    torch.set_num_threads(previous)


def gaussian(size, centre, rotated=False):
    yy, xx = np.mgrid[:size, :size]
    x, y = xx - centre[0], yy - centre[1]
    exponent = (x*x + y*y) / 18 if not rotated else (x*x + 0.7*x*y + 2*y*y) / 18
    return torch.tensor(np.exp(-exponent), dtype=torch.float32)[None, None]


@pytest.mark.parametrize('rotated', [False, True])
@pytest.mark.parametrize('kernel', [1, 3])
def test_recovers_subpixel_gaussian_peak(rotated, kernel):
    centre = (25.3, 32.7)
    maps = gaussian(64, centre, rotated)
    before = maps.clone()
    points, statuses = decode_peak_comparison(maps, [(64, 64)], 64, PeakRefinementConfig(blur_kernel=kernel))
    torch.testing.assert_close(points['argmax'], heatmaps_to_points(maps, [(64, 64)], 64), rtol=0, atol=0)
    np.testing.assert_allclose(points['dark'][0, 0].numpy(), centre, atol=0.01)
    assert statuses == [['refined']]
    torch.testing.assert_close(maps, before, rtol=0, atol=0)


def test_native_half_pixel_mapping_and_padding_exclusion():
    centre = (53.3, 61.7)
    native = gaussian(128, centre)
    native[:, :, :20] = 1000  # Very strong padding response must be ignored.
    canvas = F.interpolate(native, size=(512, 512), mode='bilinear', align_corners=False)
    points, statuses = decode_peak_comparison(canvas, [(300, 600)], 512,
                                              PeakRefinementConfig(blur_kernel=1), native)
    expected_canvas = (np.asarray(centre) + 0.5) * 4 - 0.5
    np.testing.assert_allclose(points['dark'][0, 0], expected_canvas, atol=1e-4)
    original = scale_points_to_original(points['dark'], [(300, 600)], 512)
    expected_original = ((expected_canvas - np.array([0, 128]) + 0.5) * (600 / 512)) - 0.5
    np.testing.assert_allclose(original[0, 0], expected_original, atol=1e-4)
    assert statuses == [['refined']]
    assert 'native_argmax' in points
    torch.testing.assert_close(points['argmax'], heatmaps_to_points(canvas, [(300, 600)], 512), rtol=0, atol=0)


@pytest.mark.parametrize('case', ['negative', 'flat', 'border', 'singular', 'large_offset', 'thin_content'])
def test_unstable_refinements_fall_back_to_existing_argmax(case):
    maps = gaussian(64, (24.3, 30.7))
    native, sizes, config = None, [(64, 64)], PeakRefinementConfig(blur_kernel=1)
    if case == 'negative':
        maps = -torch.ones_like(maps)
    elif case == 'flat':
        maps = torch.ones_like(maps)
    elif case == 'border':
        maps = gaussian(64, (1.3, 30.2))
    elif case == 'singular':
        maps = gaussian(64, (24.3, 30.7))[:, :, 30:31, :].expand(-1, -1, 64, -1).clone()
    elif case == 'large_offset':
        config = PeakRefinementConfig(blur_kernel=1, max_offset=0.01)
    else:
        sizes = [(1, 1000)]
        native = gaussian(16, (7.3, 8.1))
    points, statuses = decode_peak_comparison(maps, sizes, 64, config, native)
    assert statuses[0][0] != 'refined'
    torch.testing.assert_close(points['dark'], points['argmax'], rtol=0, atol=0)


def test_masked_padding_cannot_influence_smoothing():
    maps = gaussian(64, (30.3, 18.2))
    sizes = [(32, 64)]  # Content rows 16..47.
    changed = maps.clone()
    changed[:, :, :16] = 1e8
    first, status1 = decode_peak_comparison(maps, sizes, 64)
    second, status2 = decode_peak_comparison(changed, sizes, 64)
    torch.testing.assert_close(first['dark'], second['dark'], rtol=0, atol=0)
    assert status1 == status2 == [['content_boundary']]


@pytest.mark.parametrize('options', [{'blur_kernel': 2}, {'blur_sigma': -1}, {'max_offset': 0}, {'epsilon': float('nan')}])
def test_invalid_refinement_options_rejected(options):
    with pytest.raises(ValueError):
        PeakRefinementConfig(**options)


def test_nonfinite_heatmap_rejected():
    maps = gaussian(64, (32.3, 30.7))
    maps[0, 0, 1, 1] = float('nan')
    with pytest.raises(ValueError, match='finite'):
        decode_peak_comparison(maps, [(64, 64)], 64)


def test_singular_local_hessian_falls_back_before_solving():
    xs = np.arange(33) - 16
    ridge = np.tile(np.exp(-xs**2 / 18), (33, 1))
    offset, status = refine_peak(ridge, (16, 16), PeakRefinementConfig(blur_kernel=1))
    assert offset is None
    assert status == 'non_concave_or_flat_fit'


def test_capture_hook_is_removed_when_model_raises():
    class FailingModel(torch.nn.Module):
        def __init__(self):
            super().__init__()
            self.heatmap_heads = torch.nn.ModuleList([torch.nn.Identity()])

        def forward(self, inputs):
            self.heatmap_heads[-1](inputs)
            raise RuntimeError('deliberate forward failure')

    model = FailingModel()
    with pytest.raises(RuntimeError, match='deliberate'):
        forward_for_comparison(model, torch.ones(1, 1, 8, 8), True)
    assert not model.heatmap_heads[-1]._forward_hooks


@pytest.fixture
def saved_study(tmp_path):
    data = tmp_path / 'data'
    data.mkdir()
    image = np.arange(24 * 40, dtype=np.uint8).reshape(24, 40)
    cv2.imwrite(str(data / 'sample.png'), image)
    (data / 'marks.txt').write_text('sample.png (10, 8) (20, 16)\n')
    fold = data / 'folds' / 'repetition_2'
    fold.mkdir(parents=True)
    (fold / 'val_f3.txt').write_text('sample\n')
    (fold / 'training_f3.txt').write_text('another\n')
    paths = []
    for name in ('unet_basic', 'stacked_hourglass'):
        directory = tmp_path / name / 'repetition_2' / 'fold_3'
        directory.mkdir(parents=True)
        path = inference_fixtures.StandaloneInferenceTests().write_checkpoint(directory, name)
        payload = torch.load(path, weights_only=False)
        payload['metadata']['data'].update(image_data_dir=str(data), fold_lists_path=str(data / 'folds'), mark_list_file=str(data / 'marks.txt'))
        torch.save(payload, path)
        paths.append(path)
    return paths, data


def test_native_capture_preserves_final_output_and_removes_hook(saved_study):
    path = saved_study[0][1]
    loaded = load_model_from_checkpoint(path, 'cpu')
    inputs = torch.randn(1, 1, 16, 16)
    with torch.inference_mode():
        original = loaded.model(inputs)['heatmaps']
    canvas, native = forward_for_comparison(loaded.model, inputs, True)
    assert native.shape[-2:] == (4, 4)
    torch.testing.assert_close(canvas, original, rtol=0, atol=0)
    torch.testing.assert_close(canvas, F.interpolate(native, size=(16, 16), mode='bilinear', align_corners=False), rtol=0, atol=0)
    assert not loaded.model.heatmap_heads[-1]._forward_hooks


def test_saved_checkpoints_export_comparable_validation_and_paired_results(saved_study, tmp_path):
    paths, data = saved_study
    for path in paths:
        loaded = load_model_from_checkpoint(path, 'cpu')
        config = build_config_from_checkpoint_metadata(loaded.metadata, path.parent / 'existing')
        baseline = run_heatmap_inference_for_records(loaded.model, config,
            [HeatmapImageRecord('sample', data / 'sample.png', [(10, 8), (20, 16)])], 'cpu')
        reference_dir = path.parent / 'validation_best_loss'
        reference_dir.mkdir()
        write_csv_rows(reference_dir / 'validation_endpoints.csv', baseline[0]['endpoint_rows'])
    output = tmp_path / 'comparison'
    rows = run_comparison(ComparisonConfig(paths, output, expected_count=2, device='cpu',
                                           save_overlays=True, save_raw_heatmaps=True))
    assert len(rows) == 5  # Two decoders for U-Net, plus a native control for hourglass.
    assert {row['decoder'] for row in rows} == {'argmax', 'dark', 'native_argmax'}
    assert all(row['endpoint_count'] == 2 for row in rows)
    assert json.loads((output / 'experiment.json').read_text())['status'] == 'complete'
    assert len(list(output.rglob('validation_endpoints.csv'))) == 5
    assert len(list(output.rglob('*.npz'))) == 2
    for metadata_path in output.rglob('comparison_metadata.json'):
        metadata = json.loads(metadata_path.read_text())
        assert metadata['saved_argmax_audit']['status'] == 'matched'
        assert metadata['saved_argmax_audit']['endpoints_outside_tolerance'] == 0
        assert metadata['checkpoint_sha256']
        for directory in ('argmax', 'dark'):
            with (metadata_path.parent / directory / 'validation_endpoints.csv').open() as handle:
                assert all(row['dataset_split'] == 'validation' for row in csv.DictReader(handle))
            workbook = load_workbook(metadata_path.parent / directory / 'validation_summary.xlsx')
            assert workbook.sheetnames == ['validation_image_summary', 'validation_endpoints']
            workbook.close()
    with pytest.raises(ValueError, match='existing results are preserved'):
        run_comparison(ComparisonConfig(paths, output, device='cpu'))


def test_preflight_rejects_wrong_recipe_count_and_split(saved_study, tmp_path):
    paths, data = saved_study
    config = ComparisonConfig(paths, tmp_path / 'out', expected_count=30)
    with pytest.raises(ValueError, match='Expected 30'):
        prepare_jobs(config)
    config.expected_count = None
    config.required_settings = {'hourglass_stacks': 6}
    with pytest.raises(ValueError, match='REQUIRED_TRAINING_SETTINGS'):
        prepare_jobs(config)
    config.required_settings = {}
    (data / 'folds' / 'repetition_2' / 'val_f3.txt').write_text('sample\nsample\n')
    with pytest.raises(ValueError, match='duplicate'):
        prepare_jobs(config)


def test_excludes_fold_all_and_deduplicates_paths(tmp_path):
    for fold in ('1', 'all'):
        directory = tmp_path / f'fold_{fold}'
        directory.mkdir()
        (directory / 'model_best_validation_pixel_error.pth').touch()
    config = ComparisonConfig([tmp_path, tmp_path / 'fold_1'], tmp_path / 'out')
    assert discover_checkpoints(config) == [tmp_path / 'fold_1' / 'model_best_validation_pixel_error.pth']


def test_metrics_use_original_pixel_thresholds_and_expose_tail_errors():
    rows = [{'error_px': error} for error in (0, 5, 10, 20, 100)]
    result = error_summary(rows, ComparisonConfig([], 'unused'))
    assert result['within_10px_percent'] == 60
    assert result['large_error_count'] == 1
    assert result['large_error_mean_px'] == 100
    assert result['p95_error_px'] > 20
