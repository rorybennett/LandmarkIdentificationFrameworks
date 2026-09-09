"""Cross-framework output contracts; run from the repository root."""
import sys
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT / 'IPV'), str(ROOT / 'Heatmaps')]
from IPV.utils import landmark_inference_utils as ipv
from IPV.train_model import TrainModel as IPVTrainModel
from Heatmaps.utils import heatmap_inference_utils as heatmaps
from Heatmaps.utils import visualisation_utils as visuals
from Heatmaps.train_model import TrainModel as HeatmapTrainModel


class StandardisationTests(unittest.TestCase):
    def test_response_colours_repeat_and_overlays_match(self):
        image = np.full((40, 40, 3), 60, dtype=np.uint8)
        maps = np.zeros((10, 40, 40), dtype=np.float32)
        for index in range(10):
            maps[index, index + 2, index + 2] = index + 1
        maps[0, 0, 0] = -1
        self.assertEqual(ipv.POINT_COLOURS, visuals.POINT_COLOURS)
        for index in range(1, 11):
            self.assertEqual(ipv.get_point_colour(index), visuals.get_point_colour(index))
        self.assertEqual(visuals.get_point_colour(1), visuals.get_point_colour(9))
        for points in ([], [(10, 10), (20, 20)]):
            np.testing.assert_array_equal(
                ipv.create_combined_heatmap_overlay(image, maps, points),
                visuals.create_combined_heatmap_overlay(image, maps, points))
        np.testing.assert_array_equal(
            ipv.create_point_overlay(image, [(10, 10)], [(12, 12)]),
            visuals.create_point_overlay(image, [(10, 10)], [(12, 12)]))
        np.testing.assert_array_equal(
            ipv.create_combined_heatmap_overlay(image, np.zeros_like(maps), []),
            visuals.create_combined_heatmap_overlay(image, np.zeros_like(maps), []))

    def test_workbook_sheet_names_match_for_validation_and_inference(self):
        with tempfile.TemporaryDirectory() as temporary_dir:
            root = Path(temporary_dir)
            result = ipv.build_result(ipv.LandmarkImageRecord('sample', Path('sample.png'), [(0, 0)]),
                                      [(3, 4)], [(0, 0)], [1], 1, 1, 'best_validation_loss')
            inferer = object.__new__(ipv.LandmarkImageInferer)
            for label in ('validation', 'inference'):
                inferer.config = SimpleNamespace(output_dir=root, run_label=label)
                output = inferer.save_combined_summaries([result])['summary_xlsx']
                other = root / 'heatmaps.xlsx'
                if label == 'validation':
                    HeatmapTrainModel.write_validation_workbook(other, [result['summary']], result['endpoint_rows'])
                else:
                    heatmaps.write_summary_workbook(other, [result['summary']], result['endpoint_rows'])
                first, second = load_workbook(output), load_workbook(other)
                try:
                    self.assertEqual(first.sheetnames, second.sheetnames)
                finally:
                    first.close()
                    second.close()

    def test_prediction_columns_and_model_identity_match(self):
        summary = {'dataset_split': 'inference', 'repetition': 1, 'fold': 1,
                   'sample_name': 'sample', 'network_name': 'network', 'mean_error_px': 5.0}
        endpoint = {'sample_name': 'sample', 'point_index': 1, 'target_x': 0, 'target_y': 0,
                    'pred_x': 3, 'pred_y': 4, 'error_px': 5.0}
        self.assertEqual(ipv.build_prediction_rows([summary], [endpoint]),
                         heatmaps.build_prediction_rows([summary], [endpoint]))
        trainer = object.__new__(HeatmapTrainModel)
        trainer.data_config = SimpleNamespace(repetition=1, fold=1)
        trainer.model_config = SimpleNamespace(network_name='vitpose')
        rows = [trainer.create_image_summary_row('sample', 'sample.png', 40, 40, [5]),
                *trainer.create_endpoint_rows('sample', 'sample.png', [(0, 0)], [(3, 4)], [5]),
                trainer.create_prediction_row('sample', [(0, 0)], [(3, 4)], [5])]
        for row in rows:
            self.assertEqual(row['network_name'], 'vitpose')

    def test_inference_paths_are_empty_and_required_paths_are_checked(self):
        from IPV import infer_landmarks as ipv_script
        from Heatmaps import infer_landmarks as heatmap_script
        for script in (ipv_script, heatmap_script):
            for field in ('MODEL_PATH', 'INPUT_PATH', 'OUTPUT_DIR', 'GROUND_TRUTH_MARK_LIST_PATH'):
                self.assertEqual(getattr(script, field), '')
            with self.assertRaisesRegex(ValueError, 'Set MODEL_PATH'):
                script.main()
            with patch.object(script, 'MODEL_PATH', 'model.pth'):
                with self.assertRaisesRegex(ValueError, 'Set INPUT_PATH'):
                    script.main()
                with patch.object(script, 'INPUT_PATH', 'image.png'):
                    with self.assertRaisesRegex(ValueError, 'Set OUTPUT_DIR'):
                        script.main()

    def test_ipv_plot_uses_single_loss_panel_and_error_axis(self):
        trainer = object.__new__(IPVTrainModel)
        history = {'epoch': [1, 2], 'training_loss': [2, 1], 'validation_loss': [3, 2],
                   'validation_error_px': [5, 4]}
        import matplotlib.pyplot as plt
        with tempfile.TemporaryDirectory() as temporary_dir:
            output = Path(temporary_dir) / 'plot.png'
            with patch.object(trainer, 'get_plot_path', return_value=output), patch('IPV.train_model.plt.close'):
                trainer.save_history_plot(history)
                figure = plt.gcf()
            self.assertEqual(len(figure.axes), 2)
            self.assertEqual(figure.axes[0].get_ylabel(), 'Loss')
            self.assertEqual(figure.axes[1].get_ylabel(), 'Mean endpoint error (px)')
            self.assertTrue(output.is_file())
            plt.close(figure)


if __name__ == '__main__':
    unittest.main()
